"""현대 국어 사용 빈도 조사 2 xlsx -> build/freq.jsonl (02-05).

02-01 재조사(2026-09-17) 결과 원본은 csv 가 아니라 xlsx 다(00-01 예상과 다름).
난이도 산식(02-08)의 **주 지표**. 실제 파일은
`tools/raw/현대 국어 사용 빈도 조사 2/일반어휘통계.xlsx` (82,501행, 헤더 `순위,빈도,어휘,풀이,품사`).
같은 폴더의 구통계/음절통계/자모통계/조사통계/어미통계는 단어 빈도가 아니므로 무시한다
(tools/README.md "3. 현대 국어 사용 빈도 조사 2" 참조).

**어휘 컬럼에 동형어 번호가 붙어 있다** (예: '가01', '가02' -- 82,501행 중 23,912행).
02-07.merge.md "막히면" 절의 지시대로, 여기서 `hangul.normalize_headword` 를 적용한
뒤에 빈도를 결합한다 -- 그래야 `merged.py` 가 기대하는 정규화된 표제어와 매칭된다.
`normalize.py` 대신 `hangul.py` 에서 가져오는 이유는 순환 참조를 피하기 위해서다
(normalize.py 는 entries.*.jsonl 을 읽고, parse_freq.py 는 그보다 먼저/독립적으로 돈다).
"""
from pathlib import Path

import openpyxl

from . import config, io_util
from .hangul import normalize_headword

OUT = "freq.jsonl"

# README 3절 조사 결과(헤더 1행, `순위,빈도,어휘,풀이,품사`).
FREQ_COL_WORD = "어휘"
FREQ_COL_RANK = "순위"      # 순위 컬럼이 없는 원본이면 None 으로 바꾼다
FREQ_COL_COUNT = "빈도"

_MISSING = object()


def _src_files(use_fixtures: bool) -> list[Path]:
    base = config.FIXTURES if use_fixtures else config.RAW
    if use_fixtures:
        return sorted(base.glob("*freq*.xlsx")) or sorted((base / "freq").glob("*.xlsx"))
    # 실데이터 폴더명("현대 국어 사용 빈도 조사 2")은 이름 변경이 금지돼 있다(00-01).
    # 같은 폴더의 구통계.xlsx 등을 걸러내려면 "일반어휘통계" 파일명을 정확히 찾는다.
    return sorted(base.rglob("일반어휘통계*.xlsx"))


def _rows_from_workbook(path: Path):
    """xlsx 첫 시트를 `{컬럼명: 값}` dict 의 이터레이터로 돌려준다 (csv.DictReader 대응)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        header = next(rows)
        idx = {name: i for i, name in enumerate(header) if name}
        for row in rows:
            yield {name: row[i] if i < len(row) else None for name, i in idx.items()}
    finally:
        wb.close()


def _to_int(v):
    """xlsx 숫자 셀은 보통 이미 int/float 다. 문자열(천단위 구분자 포함)도 방어적으로 받는다."""
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        return int(v)
    s = str(v).replace(",", "").strip()
    return int(s) if s.lstrip("-").isdigit() else None


def _assign_ranks(rows: list[dict]) -> None:
    """순위 컬럼이 없을 때 count 내림차순으로 rank 를 직접 부여한다 (02-05 출력 스키마).

    동률은 같은 순위를 준다. 그 뒤 rank 가 건너뛰어도 무방하다 — 02-08은 정규화만 한다.
    """
    rows.sort(key=lambda r: -(r["count"] or 0))
    prev, rank = _MISSING, 0
    for i, r in enumerate(rows, 1):
        if r["count"] != prev:
            rank, prev = i, r["count"]
        r["rank"] = rank


def _best_by_headword(rows: list[dict]) -> dict[str, dict]:
    """같은 표제어가 품사별로 여러 번 나온다 (`사람/명사`, `사람/의존명사`).

    게임의 `word` 는 표제어 문자열이 PK 이므로 (DESIGN 4절) **가장 높은 순위(작은 rank)만**
    남긴다. 난이도상 보수적인(= 쉽게 보는) 선택이다 — 02-05 "중복 표제어 처리".
    """
    best: dict[str, dict] = {}
    for r in rows:
        cur = best.get(r["headword"])
        if cur is None or r["rank"] < cur["rank"]:
            best[r["headword"]] = r
    return best


def run(use_fixtures: bool = False) -> int:
    files = _src_files(use_fixtures)
    if not files:
        raise FileNotFoundError(
            "빈도 xlsx 없음. tools/raw/현대 국어 사용 빈도 조사 2/일반어휘통계.xlsx 확인 (00-01)")

    stats = {"files": len(files), "read": 0, "drop_no_word": 0}
    rows: list[dict] = []
    for p in files:
        print(f"    freq: {p.name}")
        for r in _rows_from_workbook(p):
            stats["read"] += 1
            w_raw = str(r.get(FREQ_COL_WORD) or "").strip()
            if not w_raw:
                stats["drop_no_word"] += 1
                continue
            w = normalize_headword(w_raw)
            if not w:
                stats["drop_no_word"] += 1
                continue
            cnt = _to_int(r.get(FREQ_COL_COUNT))
            rk = _to_int(r.get(FREQ_COL_RANK)) if FREQ_COL_RANK else None
            rows.append({"headword": w, "rank": rk, "count": cnt})

    if any(r["rank"] is None for r in rows):
        # 순위 없음 -> count 내림차순으로 직접 부여
        _assign_ranks(rows)

    best = _best_by_headword(rows)
    stats["drop_dup"] = len(rows) - len(best)
    stats["kept"] = io_util.write_jsonl(config.BUILD / OUT, best.values())
    print(f"    freq: {stats}")
    return stats["kept"]
